# ============================================
# Crear un archivo de Excel con openpyxl
# ============================================

from openpyxl import Workbook

# Creamos un libro de Excel nuevo (como abrir un Excel en blanco)
libro = Workbook()

# Seleccionamos la hoja activa (la primera pestaña)
hoja = libro.active

# Le ponemos un título a la hoja
hoja.title = "Inventario"

# Escribimos encabezados en la primera fila
hoja["A1"] = "Producto"
hoja["B1"] = "Cantidad"

# Guardamos el archivo con el nombre que queramos
libro.save("inventario.xlsx")

print("Archivo Excel creado con éxito")

# ============================================
# Llenar el Excel con datos usando zip y for
# ============================================

productos = ["Saltin", "Festival", "Waffer", "Tosh"]
cantidades = [50, 30, 0, 15]

# Empezamos en la fila 2, porque la fila 1 ya tiene los encabezados
fila = 2

for producto, cantidad in zip(productos, cantidades):
    hoja.cell(row=fila, column=1, value=producto)
    hoja.cell(row=fila, column=2, value=cantidad)
    fila = fila + 1

libro.save("inventario.xlsx")

print("Datos agregados con éxito")